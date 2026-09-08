#!/usr/bin/env python3
"""
Apply a Data-Team results spreadsheet to a region's data.

Reads an .xlsx with two tabs:
  - accelerator summary tab: accelerator | dealroom_profile_url | website |
      website_on_file_reliable | city | activity_status | est_startups_per_year |
      dealroom_portfolio_current | evidence | evidence_url | evidence_date | confidence
  - a per-company "results" tab: accelerator | company_name | website | ... |
      action | round_added | cohort_source | ...

For every accelerator on the summary tab it:
  * updates data/<region>_accelerators.json  (dealroom_portfolio = before + companies
    attached this run; keeps website),
  * updates the classifier batch files (website_matches_dealroom / website_status /
    evidence / evidence_url / evidence_date / cohort_throughput_estimate / confidence),
  * appends coverage-history measurement(s) marking it CHECKED now (baseline + post
    when the portfolio grew; a single "checked, no change" row otherwise) — idempotent,
  * appends per-company rows to the discovery log,
  * writes low-throughput accelerators (< THRESHOLD startups/yr) into
    data/<region>_overrides.json as MAINTAIN with no re-check schedule.

Run:  python3 scripts/apply_sheet_updates.py "<xlsx path>" --region brazil
Then: aggregate_targeting.py --region brazil ; update_tracking.py --region brazil ; build_dashboard.py
"""
import argparse, json, os, re, unicodedata, datetime
import openpyxl

BASE = os.path.join(os.path.dirname(__file__), "..")
DATA = os.path.join(BASE, "data")

ap = argparse.ArgumentParser()
ap.add_argument("xlsx")
ap.add_argument("--region", default="brazil")
ap.add_argument("--today", default=datetime.date.today().isoformat())
ap.add_argument("--baseline-date", default="2026-08-04", help="date to stamp the pre-run baseline point")
ap.add_argument("--low-throughput", type=float, default=10.0,
                help="accelerators with est_startups_per_year strictly below this move to MAINTAIN, no due date")
a = ap.parse_args()
REGION = a.region

def p(name): return os.path.join(DATA, f"{REGION}_{name}")
def load(name, default=None):
    fp = os.path.join(DATA, name)
    return json.load(open(fp)) if os.path.exists(fp) else default
def load_jsonl(name):
    fp = os.path.join(DATA, name)
    return [json.loads(l) for l in open(fp) if l.strip()] if os.path.exists(fp) else []

def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    s = re.sub(r"\(.*?\)", "", s)
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()

def cell(v): return "" if v is None else (v.strip() if isinstance(v, str) else v)

def fmt_date(v):
    if v is None or v == "": return None
    if isinstance(v, datetime.datetime): return v.strftime("%Y-%m")
    if isinstance(v, (int, float)): return str(int(v))
    return str(v).strip()

# ---------- load workbook ----------
wb = openpyxl.load_workbook(a.xlsx, data_only=True)
# summary tab = the one WITHOUT a "company_name" column; results tab has it
summary_ws = results_ws = None
for ws in wb.worksheets:
    hdr = [cell(c) for c in next(ws.iter_rows(values_only=True))]
    if "company_name" in hdr: results_ws = ws
    elif "accelerator" in hdr: summary_ws = ws
assert summary_ws and results_ws, "could not identify summary + results tabs"

def rows_of(ws):
    rs = list(ws.iter_rows(values_only=True)); H = [cell(c) for c in rs[0]]
    return [{H[i]: r[i] for i in range(len(H)) if H[i]} for r in rs[1:] if any(c is not None for c in r)]

summary = rows_of(summary_ws)
results = rows_of(results_ws)

# ---------- map accelerators -> uuid ----------
accel_json = load(f"{REGION}_accelerators.json")
url2a = {x.get("dealroom_url"): x for x in accel_json if x.get("dealroom_url")}
name2uuid = {}
recs = []
for s in summary:
    m = url2a.get(cell(s.get("dealroom_profile_url")))
    if not m:
        print("  ! no uuid match for", s.get("accelerator")); continue
    r = {
        "uuid": m["uuid"], "name": m["name"], "before_pf": m.get("dealroom_portfolio") or 0,
        "site_ok": str(cell(s.get("website_on_file_reliable"))).lower().startswith("y"),
        "est": s.get("est_startups_per_year"),
        "evidence": cell(s.get("evidence")) or None,
        "evidence_url": cell(s.get("evidence_url")) or None,
        "evidence_date": fmt_date(s.get("evidence_date")),
        "confidence": cell(s.get("confidence")) or None,
    }
    recs.append(r)
    name2uuid[norm(m["name"])] = m["uuid"]

# results-tab name aliases that don't norm-match a summary name
ALIASES = {"liga ventures oxigenio": "oxigenio aceleradora"}
def result_uuid(nm):
    n = norm(nm)
    if n in name2uuid: return name2uuid[n]
    if n in ALIASES and norm(ALIASES[n]) in name2uuid: return name2uuid[norm(ALIASES[n])]
    return None

# ---------- tally additions per uuid ----------
add = {r["uuid"]: {"new": 0, "attached": 0} for r in recs}
disc_rows = []
ACT2DISC = {"create_new": "created", "update_existing": "linked", "skip_duplicate+enhance": "linked",
            "skip_duplicate": "already_present", "skip": "already_present", "skip_no_domain": "already_present",
            "skip_curator_instruction": "already_present", "human_review": "to_add"}
for row in results:
    nm = cell(row.get("accelerator"))
    if not nm or nm == "#ERROR!": continue
    company = cell(row.get("company_name"))
    action = cell(row.get("action"))
    if company == "(cohort note)" or action == "note" or not company: continue
    u = result_uuid(nm)
    if not u:
        print("  ! results accelerator unmatched:", nm); continue
    ra = str(cell(row.get("round_added"))).lower()
    if action == "create_new": add[u]["new"] += 1
    if ra.startswith("yes"): add[u]["attached"] += 1
    disc_rows.append({
        "date": a.today, "accelerator_uuid": u, "accelerator": nm, "company": company,
        "domain": cell(row.get("website")) or None,
        "action": ACT2DISC.get(action, "to_add"),
        "cohort": cell(row.get("cohort_source")) or None,
        "dealroom_url": cell(row.get("dealroom_url_created")) or cell(row.get("dealroom_url_existing")) or None,
    })

# ---------- 1) update <region>_accelerators.json portfolio counts ----------
by_uuid = {x["uuid"]: x for x in accel_json}
for r in recs:
    after = r["before_pf"] + add[r["uuid"]]["attached"]
    by_uuid[r["uuid"]]["dealroom_portfolio"] = after
json.dump(accel_json, open(p("accelerators.json"), "w"), ensure_ascii=False, indent=2)

# ---------- 2) update classifier batch files ----------
import glob
patch = {r["uuid"]: r for r in recs}
for bf in glob.glob(p("targeting_batch_*.json")):
    arr = json.load(open(bf)); changed = False
    for row in arr:
        r = patch.get(row.get("uuid"))
        if not r: continue
        row["website_matches_dealroom"] = bool(r["site_ok"])
        row["website_status"] = "live_and_relevant" if r["site_ok"] else "needs_update"
        row["activity_status"] = "active"
        if r["evidence"]: row["evidence"] = r["evidence"]
        if r["evidence_url"]: row["evidence_url"] = r["evidence_url"]
        if r["evidence_date"]: row["evidence_date"] = r["evidence_date"]
        if r["confidence"]: row["confidence"] = r["confidence"]
        if isinstance(r["est"], (int, float)): row["cohort_throughput_estimate"] = int(r["est"])
        changed = True
    if changed:
        json.dump(arr, open(bf, "w"), ensure_ascii=False, indent=2)

# ---------- 3) append coverage-history (idempotent) ----------
hist = load_jsonl(f"{REGION}_coverage_history.jsonl")
have = {(m.get("uuid"), m.get("measured_at")) for m in hist}
new_hist = []
def rec(uuid, name, date, pf, linked, new, note):
    if (uuid, date) in have: return
    have.add((uuid, date))
    new_hist.append({"uuid": uuid, "name": name, "measured_at": date, "method": "manual",
        "website_portfolio": None, "linked": linked, "present_unlinked": None, "absent": None,
        "linkage_coverage_pct": None, "company_coverage_low_pct": None, "company_coverage_high_pct": None,
        "new_profiles_created": new, "dealroom_portfolio": pf, "note": note})
for r in recs:
    u = r["uuid"]; before = r["before_pf"]; att = add[u]["attached"]; new = add[u]["new"]; after = before + att
    if att > 0:  # portfolio grew — record a pre-run baseline so growth is visible
        rec(u, r["name"], a.baseline_date, before, 0, 0, "baseline (pre-run portfolio)")
    note = (f"Data-Team run: +{new} profiles created, {att} companies linked to portfolio ({before}→{after})."
            if (new or att) else "Checked this run; no additions.")
    rec(u, r["name"], a.today, after, att, new, note)
if new_hist:
    with open(p("coverage_history.jsonl"), "a") as f:
        for m in new_hist: f.write(json.dumps(m, ensure_ascii=False) + "\n")

# ---------- 4) append discovery log (idempotent by uuid+company+date) ----------
dl = load_jsonl(f"{REGION}_discovery_log.jsonl")
seen = {(d.get("accelerator_uuid"), d.get("company"), d.get("date")) for d in dl}
appended = 0
with open(p("discovery_log.jsonl"), "a") as f:
    for d in disc_rows:
        k = (d["accelerator_uuid"], d["company"], d["date"])
        if k in seen: continue
        seen.add(k); f.write(json.dumps(d, ensure_ascii=False) + "\n"); appended += 1

# ---------- 5) low-throughput -> MAINTAIN, no due date ----------
overrides = load(f"{REGION}_overrides.json", {}) or {}
low = []
for r in recs:
    if isinstance(r["est"], (int, float)) and r["est"] < a.low_throughput:
        overrides[r["uuid"]] = {"tier": "MAINTAIN",
            "note": f"low activity (~{int(r['est'])}/yr) — maintenance only",
            "no_due_date": True}
        low.append(f'{r["name"]} (~{int(r["est"])}/yr)')
json.dump(overrides, open(p("overrides.json"), "w"), ensure_ascii=False, indent=2)

# ---------- report ----------
print(f"[{REGION}] applied sheet: {len(recs)} accelerators updated")
print(f"  portfolio grew: " + ", ".join(f'{r["name"]} {r["before_pf"]}→{r["before_pf"]+add[r["uuid"]]["attached"]}'
      for r in recs if add[r["uuid"]]["attached"] > 0))
print(f"  new profiles created: {sum(add[u]['new'] for u in add)} across "
      f"{sum(1 for u in add if add[u]['new'])} accelerators")
print(f"  coverage-history rows appended: {len(new_hist)}; discovery rows appended: {appended}")
print(f"  low-throughput -> MAINTAIN/no-due-date: {low or 'none'}")
